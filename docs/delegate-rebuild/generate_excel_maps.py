# -*- coding: utf-8 -*-
"""Generate Excel API registry (3 cols) + Clean Architecture map for delegate_app."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
APP = "delegate_app/lib"

HEADER_FILL = PatternFill("solid", fgColor="006A64")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = "A2"
    ws.sheet_view.rightToLeft = True


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# Endpoint | file path | what it calls / does  (exactly 3 columns)
API_ROWS = [
    ("GET https://defaultdata.alsaaeidy.com/GetHaider", f"{APP}/features/auth/data/auth_remote_datasource.dart", "يجلب كتالوج المحافظات وروابط LinkDelegate"),
    ("GET {LinkDelegate}Delegates/GetDelegateLogin/asyncId={code}/", f"{APP}/features/auth/data/auth_remote_datasource.dart", "تسجيل دخول المندوب والتحقق من الرمز → حفظ الجلسة في Secure Storage"),
    ("GET {LinkDelegate}Delegates/GetDelegateTitle/{delegateId}/", f"{APP}/features/home/presentation/home_page.dart", "عرض اسم المندوب في الرئيسية (إن وُجد الاستدعاء عبر auth/home)"),
    ("GET {LinkDelegate}Customers/GetCustomersDelegateAll/{delegateId}", f"{APP}/features/customers/data/customers_remote_datasource.dart", "جلب زبائن المندوب ثم استبدال كاش SQLite عبر CustomersRepositoryImpl"),
    ("GET {LinkDelegate}CustomersPayments/GetCustomersPaymentsCustomer/{id}", f"{APP}/features/payments/data/payments_remote_datasource.dart", "سجل تسديدات زبون (قراءة)"),
    ("POST {LinkDelegate}CustomersPaymentsRequests/PostPaymentIdempotent", f"{APP}/features/payments/data/payments_remote_datasource.dart", "رفع تسديد أوفلاين بشكل idempotent — المسار المالي الحرج"),
    ("GET https://admincompany.alsaaeidy.com/sales-gw/api/mobile-updates/{appKey}", f"{APP}/core/config/app_env.dart", "بوابة التحديثات (مفتاح التطبيق delegate)"),
    ("GET {api}/admin/Customers → rating", "api-main/.../GetPaginatedCustomersHandler.cs", "يحسب اسم التقييم فقط للوحة المحاسب الرئيسي (بدون درجة)"),
    ("LOCAL CustomerRatingCalculator (C#)", "api-main/.../Customers/Common/CustomerRatingCalculator.cs", "خوارزمية الأيام/النسبة — تُرجع ممتاز/جيد/ضعيف/متلكئ/سيء"),
    ("LOCAL SQLite customers list", f"{APP}/features/customers/data/customers_local_datasource.dart", "قراءة سريعة للزبائن من الكاش المحلي للعرض"),
    ("LOCAL SecureSessionStore", f"{APP}/core/security/secure_session_store.dart", "قراءة/كتابة asyncId و delegateId و apiBase بأمان"),
]

# layer | path | role | depends_on
ARCH_ROWS = [
    ("إقلاع", f"{APP}/main.dart", "تهيئة Injection وتشغيل App", "injection, app"),
    ("إقلاع", f"{APP}/app.dart", "MaterialApp RTL والثيم", "core/theme"),
    ("إقلاع", f"{APP}/injection.dart", "تسجيل المستودعات وحالات الاستخدام", "كل الطبقات"),
    ("core", f"{APP}/core/theme/app_colors.dart", "لوحة ألوان الهوية", "—"),
    ("core", f"{APP}/core/theme/app_text_styles.dart", "أنماط الخط", "app_colors"),
    ("core", f"{APP}/core/theme/app_theme.dart", "ThemeData موحّد", "app_colors, app_text_styles"),
    ("core", f"{APP}/core/network/dio_client.dart", "Dio + HTTPS إلزامي + قياس زمن الطلب", "secure_session, app_env"),
    ("core", f"{APP}/core/network/api_exception.dart", "أخطاء عربية للواجهة", "—"),
    ("core", f"{APP}/core/security/secure_session_store.dart", "جلسة في flutter_secure_storage", "—"),
    ("core", f"{APP}/core/config/app_env.dart", "عناوين الإنتاج و requireHttps", "—"),
    ("core", f"{APP}/core/database/app_database.dart", "SQLite v2 + عمود date_sale_device", "sqflite"),
    ("core", f"{APP}/core/time/iraq_date.dart", "يوم عمل بغداد", "—"),
    ("core", f"{APP}/core/location/location_service.dart", "GPS حقيقي للتحصيل", "geolocator"),
    ("domain", f"{APP}/features/auth/domain/*", "جلسة المندوب وحالات الدخول/الخروج", "—"),
    ("data", f"{APP}/features/auth/data/*", "API الدخول + تنفيذ المستودع", "domain, dio, secure"),
    ("presentation", f"{APP}/features/auth/presentation/*", "شاشة الدخول فقط", "usecases"),
    ("domain", f"{APP}/features/customers/domain/*", "كيان الزبون + جلب محلي/مزامنة", "—"),
    ("data", f"{APP}/features/customers/data/*", "Remote/Local/Repository للزبائن", "domain, dio, sqlite"),
    ("presentation", f"{APP}/features/customers/presentation/*", "قائمة وتفاصيل الزبون (بدون تقييم)", "usecases"),
    ("domain", f"{APP}/features/payments/domain/*", "قواعد التحصيل والمزامنة", "—"),
    ("data", f"{APP}/features/payments/data/*", "طابور أوفلاين + رفع idempotent", "domain, dio, sqlite"),
    ("presentation", f"{APP}/features/payments/presentation/*", "تحصيل اليوم وورقة المبلغ", "usecases"),
    ("application", "api-main/.../Customers/Common/CustomerRatingCalculator.cs", "حاسبة اسم التقييم للمحاسب الرئيسي", "—"),
    ("application", "api-main/.../GetPaginatedCustomersHandler.cs", "يرجع rating في قائمة العملاء", "CustomerRatingCalculator"),
    ("presentation", "panel-admin-main/.../CustomerList.vue", "عمود التقييم (اسم فقط)", "GET /admin/Customers"),
    ("presentation", f"{APP}/features/home/presentation/*", "الرئيسية والشريط السفلي", "sync banner"),
    ("presentation", f"{APP}/features/sync/presentation/sync_status_banner.dart", "حالة الطابور", "payments"),
    ("presentation", f"{APP}/features/receipts/*", "PDF الإيصال", "pdf"),
]

TREE_LINES = [
    "delegate_app/",
    "  lib/",
    "    1_bootstrap → main.dart | app.dart | injection.dart",
    "    2_core/",
    "      config/ | database/ | location/ | network/ | security/ | theme/ | time/ | widgets/",
    "    3_features/",
    "      auth/ → domain | data | presentation",
    "      customers/ → domain | data | presentation",
    "      payments/ → domain | data | presentation",
    "      home/ → presentation",
    "      sync/ → presentation",
    "      receipts/ → data | presentation",
    "panel-admin-main/ → CustomerList عمود التقييم (اسم فقط)",
    "api-main/ → CustomerRatingCalculator + GetPaginatedCustomers.rating",
    "    القاعدة: presentation → usecase → repository(interface) → data impl",
]


def write_api_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "API_Registry"
    ws.append(["Endpoint", "مسار الملف المستخدم به", "ماذا يفعل / ماذا يستدعي"])
    for row in API_ROWS:
        ws.append(list(row))
    style_header(ws)
    autosize(ws, [55, 70, 70])
    out = ROOT / "Delegate_API_Documentation.xlsx"
    wb.save(out)
    return out


def write_architecture_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Architecture_Map"
    ws.append(["الطبقة", "المسار", "الدور", "يعتمد على"])
    for row in ARCH_ROWS:
        ws.append(list(row))
    style_header(ws)
    autosize(ws, [14, 75, 55, 35])

    ws2 = wb.create_sheet("Tree")
    ws2.append(["المستوى / الفرع"])
    for line in TREE_LINES:
        ws2.append([line])
    style_header(ws2)
    autosize(ws2, [90])

    ws3 = wb.create_sheet("Ratings_Feature")
    ws3.append(["الترتيب", "المسار", "الطبقة", "ملاحظة"])
    rating_detail = [
        (1, "api-main/.../Customers/Common/CustomerRatingCalculator.cs", "application", "قواعد العمل — اسم فقط"),
        (2, "api-main/.../GetPaginatedCustomersHandler.cs", "application", "يضيف rating للقائمة"),
        (3, "api-main/.../GetCustomersExcelReportHandler.cs", "application", "عمود التقييم في Excel"),
        (4, "panel-admin-main/.../CustomerList.vue", "presentation", "عمود التقييم في الجدول"),
        (5, "panel-admin-main/.../Customer/localization/ar.json", "i18n", "عنوان العمود: التقييم"),
    ]
    for r in rating_detail:
        ws3.append(list(r))
    style_header(ws3)
    autosize(ws3, [10, 75, 14, 25])

    out = ROOT / "Clean_Architecture_Map.xlsx"
    wb.save(out)
    return out


if __name__ == "__main__":
    a = write_api_excel()
    b = write_architecture_excel()
    print(f"API Excel: {a}")
    print(f"Architecture Excel: {b}")
