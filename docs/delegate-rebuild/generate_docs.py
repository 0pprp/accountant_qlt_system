# -*- coding: utf-8 -*-
"""Generate Excel API registry + Word Clean Architecture guide for delegate rebuild."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent
APP_LIB = Path(__file__).resolve().parents[2] / "delegate_collector" / "lib"

# --- API endpoints (contract for new app; compatible with SalesHaider BE_Delegate) ---
ENDPOINTS = [
    # method, path, auth, request, response, what_it_does, source_sql, client_file, mvp
    ("GET", "http://defaultdata.alsaaeidy.com/GetHaider", "لا", "—", "قائمة مدن {name, link}", "كتالوج محافظات الإنتاج", "خارجي", "features/auth/data/city_catalog_remote.dart", "نعم"),
    ("GET", "{LinkDelegate}Delegates/GetDelegateLogin/asyncId={code}/", "AsyncId في المسار (قديم)", "—", "{delegateId, asyncId, ...}", "تسجيل دخول المندوب والتحقق من الرمز", "Delegates", "features/auth/data/auth_remote_datasource.dart", "نعم"),
    ("GET", "{LinkDelegate}Delegates/GetDelegateTitle/{delegateId}/", "جلسة", "—", "عنوان/اسم المندوب", "عرض اسم المندوب في الرئيسية", "Delegates", "features/home/data/home_remote_datasource.dart", "نعم"),
    ("GET", "{LinkDelegate}Delegates/GetDelegateSelect/{delegateId}", "X-Async-Id", "—", "بيانات اختيار المندوب / القوائم", "تحديث بيانات المندوب المحلية", "Delegates", "features/customers/data/customers_remote_datasource.dart", "نعم"),
    ("GET", "{LinkDelegate}Customers/GetDateWeek", "X-Async-Id", "—", "تواريخ الأسبوع", "مبالغ إيصال الأسبوع", "DateWeek", "features/customers/data/customers_remote_datasource.dart", "نعم"),
    ("GET", "{LinkDelegate}Customers/GetCustomersDelegateAll/{delegateId}", "X-Async-Id", "—", "قائمة زبائن المندوب كاملة", "مزامنة كاش الزبائن SQLite", "Customers", "features/customers/data/customers_remote_datasource.dart", "نعم"),
    ("GET", "{LinkDelegate}Customers/GetCustomersDataInfo/{customerId}", "X-Async-Id", "—", "تفاصيل زبون", "ملف الزبون", "Customers", "features/customers/data/customers_remote_datasource.dart", "نعم"),
    ("GET", "{LinkDelegate}customers/{id}/notes?asyncId=", "AsyncId query", "—", "ملاحظات الزبون", "قراءة الملاحظات", "CustomerNotes", "features/customers/data/customers_remote_datasource.dart", "لاحقاً"),
    ("POST", "{LinkDelegate}customers/{id}/notes", "AsyncId body", "{asyncId, noteText}", "نجاح", "إضافة ملاحظة", "CustomerNotes", "features/customers/data/customers_remote_datasource.dart", "لاحقاً"),
    ("GET", "{LinkDelegate}CustomersPayments/GetCustomersPaymentsCustomer/{id}", "X-Async-Id", "—", "تسديدات الزبون", "سجل الإيصالات", "CustomersPayments", "features/payments/data/payments_remote_datasource.dart", "نعم"),
    ("GET", "{LinkDelegate}CustomersPayments/GetCustomersPaymentsCustomerName/id=&customerName=", "X-Async-Id", "—", "بحث إيصالات", "بحث بالاسم", "CustomersPayments", "features/payments/data/payments_remote_datasource.dart", "لاحقاً"),
    ("GET", "{LinkDelegate}CustomersSales/GetCustomersSalesCustomer/{id}", "جلسة", "—", "مبيعات الزبون", "سجل المبيعات", "CustomersSales", "features/customers/data/customers_remote_datasource.dart", "لاحقاً"),
    ("POST", "{LinkDelegate}CustomersPaymentsRequests/PostPaymentIdempotent", "AsyncId في الجسم", "CustomerId, DelegateId, Amount, Location, ClientPaymentId, CreatedAtUtc, ReceiptNumber, AsyncId, Device*", "200/201 أو تكرار=نجاح", "رفع تسديد أوفلاين بشكل idempotent — المسار المالي الحرج", "CustomersPaymentsRequest → صندوق عبر BE_Company", "features/payments/data/payments_remote_datasource.dart", "نعم"),
    ("GET", "{LinkDelegate}Delegates/SalesRequests?asyncId=", "AsyncId", "—", "قائمة طلبات البيع", "حالة طلبات البيع", "SalesRequests", "features/home/data/home_remote_datasource.dart", "لاحقاً"),
    ("POST", "{LinkDelegate}Delegates/SalesRequests", "AsyncId", "جسم الطلب", "نجاح", "إنشاء طلب بيع", "SalesRequests", "features/home/data/home_remote_datasource.dart", "لاحقاً"),
    ("POST", "{LinkDelegate}Delegates/FirstCollectionEvaluation", "AsyncId", "تقييم محل", "نجاح", "تقييم أول تحصيل لزبون جديد", "Evaluations", "features/payments/data/payments_remote_datasource.dart", "لاحقاً"),
    ("POST", "{LinkDelegate}delegate/complaints", "AsyncId", "نص الشكوى", "نجاح", "شكوى للمدير المفوض", "Complaints", "features/home/data/home_remote_datasource.dart", "لاحقاً"),
    ("POST", "{LinkDelegate}delegate/director-notes", "AsyncId", "ملاحظة", "نجاح", "ملاحظة للمدير", "DirectorNotes", "features/home/data/home_remote_datasource.dart", "لاحقاً"),
    ("GET", "{LinkDelegate}TrustReceipts/paged?", "X-Async-Id", "صفحات", "قائمة", "إيصالات أمانة", "TrustReceipts", "features/receipts/data/trust_receipts_remote.dart", "مرحلة 2"),
    ("POST", "{LinkDelegate}TrustReceipts", "X-Async-Id", "جسم", "نجاح", "إنشاء إيصال أمانة", "TrustReceipts", "features/receipts/data/trust_receipts_remote.dart", "مرحلة 2"),
    ("PUT", "{LinkDelegate}TrustReceipts/{id}", "X-Async-Id", "جسم", "نجاح", "تحديث إيصال أمانة", "TrustReceipts", "features/receipts/data/trust_receipts_remote.dart", "مرحلة 2"),
    ("GET", "{gateway}mobile-updates/{appKey}", "لا", "appKey=delegate-collector", "إصدار/رابط APK", "تحديث إجباري/اختياري", "Gateway", "core/update/update_remote.dart", "نعم"),
    ("GET", "{gateway}mobile-announcements/{key}", "لا", "—", "إعلانات", "بانر الرئيسية", "Gateway", "features/home/data/home_remote_datasource.dart", "لاحقاً"),
    ("GET", "{gateway}mobile-tutorial-videos/{key}", "لا", "—", "فيديوهات", "شروحات", "Gateway", "features/home/data/home_remote_datasource.dart", "لاحقاً"),
    ("GET", "{gateway}mobile-tutorial-links/{key}", "لا", "—", "روابط", "شروحات", "Gateway", "features/home/data/home_remote_datasource.dart", "لاحقاً"),
]

ACCEPTANCE = [
    ("AC-01", "دخول", "رمز صحيح + محافظة", "يفتح الرئيسية ويحفظ الجلسة في Secure Storage"),
    ("AC-02", "دخول خاطئ", "رمز خطأ", "رسالة عربية واضحة بدون حفظ جلسة"),
    ("AC-03", "تسديد أوفلاين", "بدون نت + مبلغ صالح + GPS أو رفض صريح", "صف في الطابور + قفل يوم + إيصال PDF"),
    ("AC-04", "مزامنة", "فتح النت بعد تسديدات معلّقة", "رفع PostPaymentIdempotent وحذف الصفوف الناجحة"),
    ("AC-05", "Idempotent", "إعادة رفع نفس ClientPaymentId", "لا تكرار مالي؛ الجهاز يمسح الطابور"),
    ("AC-06", "نطاق زبون", "محاولة تسديد زبون غير تابع", "رفض محلي/سيرفر"),
    ("AC-07", "GPS", "رفض صلاحية الموقع", "لا إحداثيات 1111,1111 — رسالة أو منع"),
    ("AC-08", "خروج مع طابور", "يوجد pending", "يمنع مسح الطابور أو يحذّر بقوة"),
    ("AC-09", "HTTPS إنتاج", "build release", "لا cleartext"),
    ("AC-10", "سرعة قائمة", "بعد أول مزامنة", "فتح الزبائن من SQLite أقل من ثانيتين"),
]

ARCHITECTURE_FILES = [
    # path relative to lib/, layer, arabic description
    ("main.dart", "إقلاع", "نقطة تشغيل التطبيق: تهيئة DI، Crashlytics، خدمات المزامنة، ثم تشغيل App."),
    ("app.dart", "إقلاع", "MaterialApp RTL، الثيم، والمسارات الرئيسية."),
    ("injection.dart", "إقلاع", "تسجيل الاعتماديات (GetIt): Dio، قواعد البيانات، المستودعات، حالات الاستخدام."),
    ("core/theme/app_theme.dart", "core", "ثيم موحّد: ألوان الهوية، خط Cairo، أزرار/بطاقات."),
    ("core/theme/app_colors.dart", "core", "ثوابت الألوان (لا ألوان عشوائية داخل الشاشات)."),
    ("core/network/dio_client.dart", "core", "عميل Dio مع interceptor لـ X-Async-Id وقياس الزمن ومنع HTTP في الإنتاج."),
    ("core/network/api_exception.dart", "core", "ترجمة أخطاء الشبكة لرسائل عربية للمستخدم."),
    ("core/security/secure_session_store.dart", "core", "حفظ AsyncId وLinkDelegate وDelegateId في flutter_secure_storage فقط."),
    ("core/security/url_policy.dart", "core", "سياسة المضيفين الموثوقين ورفض cleartext في release."),
    ("core/location/location_service.dart", "core", "جلب GPS الحقيقي أو فشل صريح — بدون إحداثيات وهمية."),
    ("core/time/iraq_date.dart", "core", "يوم العمل بتوقيت بغداد (بداية اليوم 03:00) وقفل التحصيل اليومي."),
    ("core/database/app_database.dart", "core", "فتح SQLite وترحيل الجداول (Customer, PendingPayment, BusinessDayLock)."),
    ("core/database/tables.dart", "core", "تعريف جداول الأوفلاين والحقول المالية الحرجة."),
    ("core/update/update_remote.dart", "core", "فحص إصدار التطبيق من الـ gateway."),
    ("core/update/update_gate.dart", "core", "شاشة تحديث إجباري قبل الدخول للعمل."),
    ("core/widgets/app_scaffold.dart", "core", "هيكل شاشة مشترك مع حالات تحميل/خطأ/فارغ."),
    ("core/widgets/empty_state.dart", "core", "حالة فارغة موحّدة."),
    ("core/widgets/error_state.dart", "core", "حالة خطأ مع زر إعادة محاولة."),
    ("core/widgets/loading_view.dart", "core", "مؤشر تحميل موحّد."),
    ("features/auth/domain/entities/delegate_session.dart", "domain", "كيان الجلسة: delegateId, asyncId, apiBase."),
    ("features/auth/domain/entities/city.dart", "domain", "كيان المحافظة/المدينة ورابط الفرع."),
    ("features/auth/domain/repositories/auth_repository.dart", "domain", "واجهة مجردة لتسجيل الدخول والخروج والتحقق من الجلسة."),
    ("features/auth/domain/usecases/login_usecase.dart", "domain", "حالة استخدام: اختيار مدينة + رمز → جلسة آمنة."),
    ("features/auth/domain/usecases/logout_usecase.dart", "domain", "خروج مع حماية طابور التسديدات المعلّق."),
    ("features/auth/data/auth_remote_datasource.dart", "data", "استدعاء GetDelegateLogin وكتالوج المدن."),
    ("features/auth/data/auth_repository_impl.dart", "data", "تنفيذ المستودع: remote + SecureSessionStore."),
    ("features/auth/presentation/login_page.dart", "presentation", "شاشة الدخول RTL."),
    ("features/auth/presentation/login_controller.dart", "presentation", "منطق واجهة الدخول (تحميل/خطأ)."),
    ("features/customers/domain/entities/customer.dart", "domain", "كيان الزبون للحقل والتحصيل."),
    ("features/customers/domain/repositories/customers_repository.dart", "domain", "واجهة جلب/تحديث الزبائن محلياً وعن بُعد."),
    ("features/customers/domain/usecases/refresh_customers_usecase.dart", "domain", "تحديث صامت للزبائن بدون تجميد الواجهة."),
    ("features/customers/domain/usecases/get_local_customers_usecase.dart", "domain", "قراءة سريعة من SQLite للعرض."),
    ("features/customers/data/customers_remote_datasource.dart", "data", "GetCustomersDelegateAll وGetCustomersDataInfo وGetDateWeek."),
    ("features/customers/data/customers_local_datasource.dart", "data", "CRUD زبائن في SQLite."),
    ("features/customers/data/customers_repository_impl.dart", "data", "دمج remote/local مع تحديث تفاضلي إن أمكن."),
    ("features/customers/presentation/customers_page.dart", "presentation", "قائمة الزبائن مع بحث."),
    ("features/customers/presentation/customer_detail_page.dart", "presentation", "تفاصيل الزبون وزر التسديد."),
    ("features/payments/domain/entities/pending_payment.dart", "domain", "كيان تسديد معلّق: ClientPaymentId, مبلغ, حالة مزامنة."),
    ("features/payments/domain/repositories/payments_repository.dart", "domain", "واجهة إضافة تسديد محلي ومزامنة الرفع."),
    ("features/payments/domain/usecases/collect_payment_usecase.dart", "domain", "تحصيل أوفلاين: GPS + قفل يوم + إدراج طابور + إيصال."),
    ("features/payments/domain/usecases/sync_pending_payments_usecase.dart", "domain", "رفع الطابور عبر PostPaymentIdempotent."),
    ("features/payments/data/payments_remote_datasource.dart", "data", "PostPaymentIdempotent وسجل التسديدات."),
    ("features/payments/data/payments_local_datasource.dart", "data", "طابور SQLite + قفل يوم العمل."),
    ("features/payments/data/payments_repository_impl.dart", "data", "تنفيذ المستودع المالي."),
    ("features/payments/data/payment_sync_service.dart", "data", "مزامنة خلفية عند عودة النت وresume."),
    ("features/payments/presentation/today_payments_page.dart", "presentation", "تسديدات اليوم."),
    ("features/payments/presentation/collect_payment_sheet.dart", "presentation", "إدخال المبلغ وتأكيد التحصيل."),
    ("features/receipts/domain/usecases/build_receipt_pdf_usecase.dart", "domain", "توليد PDF الإيصال."),
    ("features/receipts/data/receipt_pdf_builder.dart", "data", "بناء PDF بخط عربي."),
    ("features/receipts/presentation/receipt_preview_page.dart", "presentation", "معاينة/مشاركة الإيصال."),
    ("features/home/domain/repositories/home_repository.dart", "domain", "عنوان المندوب وبيانات الرئيسية."),
    ("features/home/data/home_remote_datasource.dart", "data", "GetDelegateTitle وإعلانات (لاحقاً)."),
    ("features/home/data/home_repository_impl.dart", "data", "تنفيذ مستودع الرئيسية."),
    ("features/home/presentation/home_shell.dart", "presentation", "الشريط السفلي: رئيسية / تسديدات / زبائن."),
    ("features/home/presentation/home_page.dart", "presentation", "لوحة الرئيسية مع حالة المزامنة."),
    ("features/sync/presentation/sync_status_banner.dart", "presentation", "عرض عدد المعلّق وزر إرسال الآن."),
]


def style_header(ws, row=1):
    fill = PatternFill("solid", fgColor="0F766E")
    font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def write_excel():
    wb = Workbook()

    ws = wb.active
    ws.title = "API_Endpoints"
    headers = [
        "#",
        "Method",
        "Path",
        "Auth",
        "Request",
        "Response",
        "ماذا يفعل",
        "مصدر SQL / النظام",
        "ملف العميل (Clean Arch)",
        "MVP؟",
    ]
    ws.append(headers)
    for i, row in enumerate(ENDPOINTS, 1):
        ws.append([i, *row])
    style_header(ws)
    widths = [5, 8, 55, 22, 35, 30, 45, 35, 50, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = w
    # fix wide columns properly
    from openpyxl.utils import get_column_letter

    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.sheet_view.rightToLeft = True

    ws2 = wb.create_sheet("Acceptance_Matrix")
    ws2.append(["ID", "السيناريو", "المدخلات", "النتيجة المتوقعة"])
    for row in ACCEPTANCE:
        ws2.append(list(row))
    style_header(ws2)
    for i, w in enumerate([10, 18, 40, 55], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"
    ws2.sheet_view.rightToLeft = True

    ws3 = wb.create_sheet("Architecture_Files")
    ws3.append(["#", "المسار تحت lib/", "الطبقة", "ماذا يعمل هذا الملف"])
    for i, (path, layer, desc) in enumerate(ARCHITECTURE_FILES, 1):
        ws3.append([i, path, layer, desc])
    style_header(ws3)
    for i, w in enumerate([5, 55, 14, 70], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = "A2"
    ws3.sheet_view.rightToLeft = True

    ws4 = wb.create_sheet("Scope")
    ws4.append(["النوع", "البند"])
    for item in [
        ("يستقبل", "كتالوج المدن GetHaider (أو demo ثابت)"),
        ("يستقبل", "BE_Delegate عبر LinkDelegate: دخول، زبائن، تسديدات، ملاحظات، طلبات"),
        ("يستقبل", "Gateway: تحديثات APK وإعلانات"),
        ("يستقبل", "الجهاز: GPS الحقيقي، SQLite، PDF"),
        ("لا يستقبل", "زبائن مناديب آخرين"),
        ("لا يستقبل", "APIs المتابع / إدارة الفرع / الصندوق مباشرة"),
        ("لا يستقبل", "إنشاء مستخدمين أو مهام المدير"),
        ("مرجع فقط", "SalesHaider delegate_application — لا نعدّله هنا"),
    ]:
        ws4.append(list(item))
    style_header(ws4)
    ws4.column_dimensions["A"].width = 14
    ws4.column_dimensions["B"].width = 70
    ws4.sheet_view.rightToLeft = True

    out = ROOT / "Delegate_API_Documentation.xlsx"
    wb.save(out)
    return out


def write_word():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    title = doc.add_heading("دليل Clean Architecture — تطبيق المندوب (delegate_collector)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(
        "هذا الملف يشرح كل ملف في المشروع الجديد ومساره ووظيفته. "
        "المرجع السلوكي من SalesHaider، والتنفيذ هنا داخل accountant_qlt_system/delegate_collector."
    )

    doc.add_heading("1) الطبقات", level=1)
    for line in [
        "presentation: الشاشات والتحكم بالواجهة فقط — لا HTTP مباشر.",
        "domain: كيانات + واجهات مستودعات + حالات استخدام (قواعد العمل).",
        "data: تنفيذ الـ API وSQLite وتنفيذ المستودعات.",
        "core: شبكة، أمان، ثيم، قاعدة بيانات مشتركة، موقع، وقت بغداد.",
    ]:
        para = doc.add_paragraph(line, style="List Bullet")
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_heading("2) تدفق التسديد المالي (لا يُكسر)", level=1)
    flow = doc.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    flow.add_run(
        "واجهة التحصيل → CollectPaymentUseCase → طابور SQLite (ClientPaymentId UUID) "
        "→ PaymentSyncService → POST PostPaymentIdempotent → حذف الصف عند النجاح "
        "→ ترحيل الصندوق على السيرفر عبر BE_Company (خارج التطبيق)."
    )

    doc.add_heading("3) فهرس الملفات", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "المسار"
    hdr[1].text = "الطبقة"
    hdr[2].text = "ماذا يعمل"
    for path, layer, desc in ARCHITECTURE_FILES:
        row = table.add_row().cells
        row[0].text = path
        row[1].text = layer
        row[2].text = desc

    doc.add_heading("4) قواعد أمان في هذه النسخة", level=1)
    for line in [
        "AsyncId فقط في flutter_secure_storage — ليس SharedPreferences.",
        "إرسال الهوية عبر header X-Async-Id قدر الإمكان؛ تجنب وضع الرمز في URL في الكود الجديد.",
        "لا GPS وهمي (1111,1111).",
        "لا cleartext في build الإنتاج.",
        "رفض تسديد زبون خارج نطاق المندوب قبل الإدراج في الطابور.",
    ]:
        para = doc.add_paragraph(line, style="List Bullet")
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_heading("5) علاقة الملفات بـ Excel", level=1)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    note.add_run(
        "ملف Excel المرافق Delegate_API_Documentation.xlsx يحتوي: "
        "ورقة API_Endpoints (كل endpoint ومسار الملف الذي يستدعيه)، "
        "ورقة Acceptance_Matrix (اختبارات القبول)، "
        "ورقة Architecture_Files (نفس فهرس هذا المستند)، "
        "ورقة Scope (من أين يستقبل / لا يستقبل)."
    )

    out = ROOT / "Clean_Architecture_File_Guide.docx"
    doc.save(out)
    return out


def write_per_file_words():
    """One short Word doc per architecture file for the user to browse."""
    out_dir = ROOT / "architecture" / "files"
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for path, layer, desc in ARCHITECTURE_FILES:
        safe = path.replace("/", "__").replace(".", "_")
        doc = Document()
        h = doc.add_heading(path, 0)
        h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for label, val in [
            ("المسار الكامل المتوقع", f"delegate_collector/lib/{path}"),
            ("الطبقة", layer),
            ("ماذا يعمل", desc),
            ("ملاحظة", "لا تستدعِ HTTP من presentation — مرّ عبر domain use case ثم data."),
        ]:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(f"{label}: ")
            run.bold = True
            p.add_run(val)
        out = out_dir / f"{safe}.docx"
        doc.save(out)
        paths.append(out)
    # index
    idx = Document()
    idx.add_heading("فهرس ملفات Word — Clean Architecture", 0).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for path, layer, desc in ARCHITECTURE_FILES:
        safe = path.replace("/", "__").replace(".", "_")
        p = idx.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run(f"{path} → architecture/files/{safe}.docx — [{layer}] {desc}")
    index_path = out_dir / "00_INDEX.docx"
    idx.save(index_path)
    return index_path, len(paths)


if __name__ == "__main__":
    x = write_excel()
    w = write_word()
    index, n = write_per_file_words()
    print(f"Excel: {x}")
    print(f"Word guide: {w}")
    print(f"Per-file docs: {n} + {index}")
